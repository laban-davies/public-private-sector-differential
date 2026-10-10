# NOTE: this script was written with AI assistance (Claude, Anthropic). See README for details.
"""
Builds the write-up figures and tables from results.csv.

results.csv is produced by parse_log.py from Stata's key_results.log,
so the route is: Master_do_file.do -> key_results.log -> parse_log.py -> results.csv -> make_outputs.py

Usage:  python make_outputs.py            (reads results.csv, writes into ./outputs)
Needs:  matplotlib, python-docx
"""
import csv
import os

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = "outputs"
os.makedirs(OUT, exist_ok=True)

YEARS = ["2021-22", "2022-23"]
COLOURS = {"2021-22": "#2a78d6", "2022-23": "#eb6834"}  # validated colourblind-safe pair
MARKERS = {"2021-22": "o", "2022-23": "D"}               # shape as a second cue, not colour alone
INK, MUTED, GRID = "#000000", "#595959", "#e6e6e6"  # black text to match the Word document

rows = list(csv.DictReader(open("results.csv")))
for r in rows:
    for k in ("est", "se", "p", "lb", "ub"):
        r[k] = float(r[k])
    r["N"] = int(r["N"])
sample = {r["fy"]: r for r in csv.DictReader(open("sample.csv"))}


def get(section):
    """{label: {year: row}} in the order labels first appear."""
    out = {}
    for r in rows:
        if r["section"] == section:
            out.setdefault(r["label"], {})[r["fy"]] = r
    return out


# ----------------------------------------------------------------- figures
plt.rcParams.update({
    # Times New Roman to match the write-up (Liberation Serif is its metric-identical fallback)
    "font.family": "serif", "font.serif": ["Times New Roman", "Liberation Serif"], "font.size": 11,
    "axes.edgecolor": MUTED, "axes.labelcolor": INK,
    "xtick.color": INK, "ytick.color": INK,
})


NOTE = ("Source: Labour Force Survey, author's calculations. Conditional estimates control for age, experience, education, "
        "region and quarter, interacted with sex. Positive = public sector pays more.")


def dot_whisker(section, filename, title, order=None, width=7.5):
    data = get(section)
    labels = order or list(data)
    n = len(labels)
    fig, ax = plt.subplots(figsize=(width, 0.42 * n + 1.5), dpi=200)
    offset = {"2021-22": 0.17, "2022-23": -0.17}
    for fy in YEARS:
        ys, xs, lo, hi = [], [], [], []
        for i, lab in enumerate(labels):
            r = data[lab][fy]
            ys.append(n - 1 - i + offset[fy])
            xs.append(r["est"]); lo.append(r["est"] - r["lb"]); hi.append(r["ub"] - r["est"])
        ax.errorbar(xs, ys, xerr=[lo, hi], fmt=MARKERS[fy], color=COLOURS[fy], ecolor=COLOURS[fy],
                    elinewidth=1.6, capsize=3, markersize=6, markeredgecolor="white",
                    markeredgewidth=1, label=fy, zorder=3)
    ax.axvline(0, color=MUTED, lw=1, ls="--", zorder=1)
    ax.set_yticks(range(n)); ax.set_yticklabels(labels[::-1])
    ax.set_ylim(-0.7, n - 0.3)
    ax.grid(axis="x", color=GRID, lw=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_xlabel("Public–private hourly pay differential (%), with 95% CI")
    h = 0.42 * n + 1.5
    top = 1 - 0.75 / h  # reserve a fixed band for title + legend whatever the chart height
    fig.text(0.01, 1 - 0.18 / h, title, fontsize=12.5, color=INK, fontweight="bold", va="top")
    handles, labs = ax.get_legend_handles_labels()
    fig.legend(handles, labs, ncol=2, frameon=False, loc="upper left", bbox_to_anchor=(0.005, 1 - 0.42 / h),
               handletextpad=0.3, fontsize=10.5)
    fig.text(0.01, 0.005, NOTE, fontsize=8, color=MUTED, wrap=True)
    fig.tight_layout(rect=(0, 0.045, 1, top))
    fig.savefig(os.path.join(OUT, filename), facecolor="white")
    plt.close(fig)


# regions sorted by the average of the two years so the pattern reads top to bottom
reg = get("Region")
reg_order = sorted(reg, key=lambda l: -(reg[l]["2021-22"]["est"] + reg[l]["2022-23"]["est"]))

# headline and sex in one chart
for r in rows:
    if r["section"] == "Furlough check":
        r["label"] = {"Oct-Mar quarters": "Oct–Mar quarters (no furlough)",
                      "Apr-Sep quarters": "Apr–Sep quarters (furlough in 2021-22)"}.get(r["label"], r["label"])
for r in rows:
    if r["section"] == "Sex":
        r["section"] = "Headline"
        r["label"] = f"Conditional – {r['label'].lower()}"

SENS = "Conditional (HIQUL15D education controls)"  # sensitivity check: reported in Table 1, kept off the chart
fig1_order = [l for l in get("Headline") if l != SENS]
dot_whisker("Headline", "fig1_headline.png", "Public–private pay differential: raw, conditional and by sex",
            order=fig1_order)
dot_whisker("Region", "fig2_region.png", "Conditional public–private pay differential by region", order=reg_order)
dot_whisker("Education", "fig3_education.png", "Conditional public–private pay differential by highest qualification", width=8.5)
dot_whisker("Furlough check", "fig4_furlough.png",
            "Furlough check: conditional differential by part of year", width=7.5)

# ------------------------------------------------------------------ tables
def cell(r):
    stars = "***" if r["p"] < 0.01 else "**" if r["p"] < 0.05 else "*" if r["p"] < 0.1 else ""
    return f"{r['est']:.1f}{stars}", f"[{r['lb']:.1f}, {r['ub']:.1f}]"


from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SIG = 0.05  # highlight threshold


def shade(cell, fill="FFF200"):
    """Fill a table cell with yellow (used for differentials significant at the 5% level)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def pval(p):
    # nlcom reports p to 3 decimal places, so 0.000 means p < 0.001
    return "<0.001" if p < 0.001 else f"{p:.3f}"


doc = Document()
# landscape so the nine-column tables fit
from docx.enum.section import WD_ORIENT
from docx.shared import Cm
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
sec.left_margin = sec.right_margin = Cm(1.8)
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(10)


def table(title, section, order=None, note=None):
    data = get(section)
    labels = order or list(data)
    doc.add_heading(title, level=2)
    t = doc.add_table(rows=1, cols=9)
    t.style = "Light List"
    hdr = ["", "2021-22", "95% CI", "p", "N", "2022-23", "95% CI", "p", "N"]
    for j, h in enumerate(hdr):
        t.rows[0].cells[j].text = h
    for lab in labels:
        c = t.add_row().cells
        c[0].text = lab
        a, b = data[lab]["2021-22"], data[lab]["2022-23"]
        c[1].text, c[2].text = cell(a)
        c[3].text = pval(a["p"])
        c[4].text = f"{a['N']:,}"
        c[5].text, c[6].text = cell(b)
        c[7].text = pval(b["p"])
        c[8].text = f"{b['N']:,}"
        for r_, cols in ((a, (1, 2, 3)), (b, (5, 6, 7))):
            if r_["p"] < SIG:
                for j in cols:
                    shade(c[j])
    for row in t.rows:
        for j in range(1, 9):
            row.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph(note or "")
    p.add_run("Percentage differential, 100×(exp(b)−1), weighted by PIWT22 with robust standard errors. "
              "p-values test whether each differential differs from zero; yellow = significant at the 5% level. "
              "*** p<0.01, ** p<0.05, * p<0.1.").italic = True


doc.add_heading("Public–private sector pay differential: results tables", level=1)

doc.add_heading("Sample", level=2)
t = doc.add_table(rows=1, cols=4)
t.style = "Light List"
for j, h in enumerate(["", "Private sector", "Public sector", "Total"]):
    t.rows[0].cells[j].text = h
for fy in YEARS:
    c = t.add_row().cells
    c[0].text = fy
    for j, k in enumerate(("private", "public", "total"), start=1):
        c[j].text = f"{int(sample[fy][k]):,}"
        c[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
doc.add_paragraph().add_run("Unweighted respondents: employees aged 20–59 in LFS waves 1 and 5 with valid hourly pay, "
                            "after trimming.").italic = True

table("Table 1. Headline estimates", "Headline",
      note="The HIQUL15D row is a sensitivity check: degree and higher degree are combined in the education controls. ")
table("Table 2. Furlough check (conditional differential)", "Furlough check",
      note="2021-22: Apr–Sep quarters were under the furlough scheme. 2022-23 is a placebo using the same calendar quarters. ")
table("Table A1. By region", "Region", order=reg_order)
table("Table A2. By highest qualification", "Education")
doc.save(os.path.join(OUT, "results_tables.docx"))

print("Written:", sorted(os.listdir(OUT)))
