# NOTE: this script was written with AI assistance (Claude, Anthropic). See README for details.
"""
Builds figures_editable.docx: native Word charts (not pictures) from results.csv.

Each chart is a real Office chart with its own embedded spreadsheet, so in Word you can
right-click > Edit Data, change colours, fonts, titles, axis ranges, etc.

Form: horizontal clustered bars (one bar per year) with custom 95% CI error bars.
(Word has no native dot-and-whisker chart with category labels; bars + error bars is the
closest fully editable equivalent.)

Usage:  python make_word_charts.py      (reads results.csv, writes outputs/figures_editable.docx)
Needs:  python-docx, openpyxl, lxml
"""
import csv
import io
import os
from xml.sax.saxutils import escape

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor

FONT = "Times New Roman"
COLOURS = {"2021-22": "2A78D6", "2022-23": "EB6834"}  # same colourblind-checked pair as the PNG charts
YEARS = ["2021-22", "2022-23"]
CHART_CT = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOTE = ("Source: Labour Force Survey, author's calculations. Conditional estimates control for age, experience, "
        "education, region and quarter, interacted with sex. Bars show the percentage differential; lines show "
        "95% confidence intervals. Positive = public sector pays more.")

# ------------------------------------------------------------------ data
rows = list(csv.DictReader(open("results.csv")))
for r in rows:
    for k in ("est", "lb", "ub"):
        r[k] = float(r[k])


def get(section):
    out = {}
    for r in rows:
        if r["section"] == section:
            out.setdefault(r["label"], {})[r["fy"]] = r
    return out


furlough_names = {"Oct-Mar quarters": "Oct–Mar quarters (no furlough)",
                  "Apr-Sep quarters": "Apr–Sep quarters (furlough in 2021-22)"}
for r in rows:
    if r["section"] == "Furlough check":
        r["label"] = furlough_names.get(r["label"], r["label"])
    if r["section"] == "Sex":
        r["section"], r["label"] = "Headline", f"Conditional – {r['label'].lower()}"

reg = get("Region")
reg_order = sorted(reg, key=lambda l: -(reg[l]["2021-22"]["est"] + reg[l]["2022-23"]["est"]))
SENS = "Conditional (HIQUL15D education controls)"

CHARTS = [
    ("Public–private pay differential: raw, conditional and by sex", "Headline",
     [l for l in get("Headline") if l != SENS]),
    ("Conditional public–private pay differential by region", "Region", reg_order),
    ("Conditional public–private pay differential by highest qualification", "Education", list(get("Education"))),
    ("Furlough check: conditional differential by part of year", "Furlough check", list(get("Furlough check"))),
]


# ------------------------------------------------------------------ embedded workbook
def workbook(labels, data):
    """Sheet1: A label | B,C estimates | D-G lower/upper CI | H-K error-bar lengths (formulas)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Group", "2021-22", "2022-23", "2021-22 lower CI", "2021-22 upper CI",
               "2022-23 lower CI", "2022-23 upper CI",
               "2021-22 minus", "2021-22 plus", "2022-23 minus", "2022-23 plus"])
    for i, lab in enumerate(labels, start=2):
        a, b = data[lab]["2021-22"], data[lab]["2022-23"]
        ws.append([lab, round(a["est"], 3), round(b["est"], 3), round(a["lb"], 3), round(a["ub"], 3),
                   round(b["lb"], 3), round(b["ub"], 3),
                   f"=B{i}-D{i}", f"=E{i}-B{i}", f"=C{i}-F{i}", f"=G{i}-C{i}"])
    ws.column_dimensions["A"].width = 42
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ chart XML
def txpr(size=1000, bold=False, colour="000000"):
    b = ' b="1"' if bold else ' b="0"'
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{size}"{b}>'
            f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="en-GB"/></a:p></c:txPr>')


def rich(text, size, bold):
    b = "1" if bold else "0"
    return (f'<c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{size}" b="{b}">'
            f'<a:latin typeface="{FONT}"/></a:defRPr></a:pPr><a:r><a:rPr lang="en-GB" sz="{size}" b="{b}">'
            f'<a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:latin typeface="{FONT}"/></a:rPr>'
            f'<a:t>{escape(text)}</a:t></a:r></a:p></c:rich></c:tx>')


def numref(ref, values):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v:.4f}</c:v></c:pt>' for i, v in enumerate(values))
    return (f'<c:numRef><c:f>{ref}</c:f><c:numCache><c:formatCode>General</c:formatCode>'
            f'<c:ptCount val="{len(values)}"/>{pts}</c:numCache></c:numRef>')


def chart_xml(title, labels, data, xlsx_rid):
    n = len(labels)
    last = n + 1
    cat_pts = "".join(f'<c:pt idx="{i}"><c:v>{escape(l)}</c:v></c:pt>' for i, l in enumerate(labels))
    cat = (f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${last}</c:f><c:strCache><c:ptCount val="{n}"/>'
           f'{cat_pts}</c:strCache></c:strRef></c:cat>')
    series = ""
    cols = {"2021-22": ("B", "H", "I"), "2022-23": ("C", "J", "K")}
    for i, fy in enumerate(YEARS):
        vcol, mcol, pcol = cols[fy]
        est = [data[l][fy]["est"] for l in labels]
        minus = [data[l][fy]["est"] - data[l][fy]["lb"] for l in labels]
        plus = [data[l][fy]["ub"] - data[l][fy]["est"] for l in labels]
        series += (
            f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${vcol}$1</c:f><c:strCache><c:ptCount val="1"/>'
            f'<c:pt idx="0"><c:v>{fy}</c:v></c:pt></c:strCache></c:strRef></c:tx>'
            f'<c:spPr><a:solidFill><a:srgbClr val="{COLOURS[fy]}"/></a:solidFill></c:spPr>'
            f'<c:invertIfNegative val="0"/>'
            f'<c:errBars><c:errBarType val="both"/><c:errValType val="cust"/><c:noEndCap val="0"/>'
            f'<c:plus>{numref(f"Sheet1!${pcol}$2:${pcol}${last}", plus)}</c:plus>'
            f'<c:minus>{numref(f"Sheet1!${mcol}$2:${mcol}${last}", minus)}</c:minus>'
            f'<c:spPr><a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></c:spPr>'
            f'</c:errBars>'
            f'{cat}<c:val>{numref(f"Sheet1!${vcol}$2:${vcol}${last}", est)}</c:val></c:ser>')
    grid = '<c:spPr><a:ln w="6350"><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill></a:ln></c:spPr>'
    axis_line = '<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="595959"/></a:solidFill></a:ln></c:spPr>'
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<c:date1904 val="0"/><c:lang val="en-GB"/><c:roundedCorners val="0"/>'
        '<c:chart>'
        f'<c:title>{rich(title, 1200, True)}<c:overlay val="0"/></c:title>'
        '<c:autoTitleDeleted val="0"/>'
        '<c:plotArea><c:layout/>'
        '<c:barChart><c:barDir val="bar"/><c:grouping val="clustered"/><c:varyColors val="0"/>'
        f'{series}<c:gapWidth val="60"/><c:overlap val="0"/>'
        '<c:axId val="111"/><c:axId val="222"/></c:barChart>'
        # category axis: first group at the top, labels kept at the far left (not at zero)
        '<c:catAx><c:axId val="111"/><c:scaling><c:orientation val="maxMin"/></c:scaling>'
        '<c:delete val="0"/><c:axPos val="l"/><c:numFmt formatCode="General" sourceLinked="0"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="low"/>'
        f'{axis_line}{txpr(1000)}'
        '<c:crossAx val="222"/><c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/>'
        '<c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>'
        # value axis: sits at the bottom of the reversed category axis
        '<c:valAx><c:axId val="222"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
        f'<c:delete val="0"/><c:axPos val="b"/><c:majorGridlines>{grid}</c:majorGridlines>'
        f'<c:title>{rich("Public–private hourly pay differential (%)", 1000, False)}<c:overlay val="0"/></c:title>'
        '<c:numFmt formatCode="0" sourceLinked="0"/><c:majorTickMark val="out"/><c:minorTickMark val="none"/>'
        f'<c:tickLblPos val="nextTo"/>{axis_line}{txpr(1000)}'
        '<c:crossAx val="111"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>'
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
        '</c:plotArea>'
        f'<c:legend><c:legendPos val="t"/><c:overlay val="0"/>{txpr(1000)}</c:legend>'
        '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/>'
        '</c:chart>'
        '<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
        f'{txpr(1000)}'
        f'<c:externalData r:id="{xlsx_rid}"><c:autoUpdate val="0"/></c:externalData>'
        '</c:chartSpace>').encode("utf-8")


# ------------------------------------------------------------------ document
doc = Document()
doc.styles["Normal"].font.name = FONT
doc.styles["Normal"].font.size = Pt(12)
package = doc.part.package
EMU_PER_CM = 360000

for k, (title, section, labels) in enumerate(CHARTS, start=1):
    data = get(section)
    xlsx_part = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{k}.xlsx"), XLSX_CT,
                     workbook(labels, data), package)
    chart_part = Part(PackURI(f"/word/charts/chart{k}.xml"), CHART_CT, b"", package)
    xlsx_rid = chart_part.relate_to(xlsx_part, RT.PACKAGE)
    chart_part._blob = chart_xml(title, labels, data, xlsx_rid)
    chart_rid = doc.part.relate_to(chart_part, RT.CHART)

    width = int(16 * EMU_PER_CM)
    height = int((3.2 + 0.9 * len(labels)) * EMU_PER_CM)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        f'<wp:extent cx="{width}" cy="{height}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{k}" name="Chart {k}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{chart_rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'))
    note = doc.add_paragraph()
    run = note.add_run(NOTE)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    doc.add_paragraph()

os.makedirs("outputs", exist_ok=True)
doc.save("outputs/figures_editable.docx")
print("Written outputs/figures_editable.docx with", len(CHARTS), "native charts")
